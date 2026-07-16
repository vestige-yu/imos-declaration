from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from .text_utils import normalize_header, safe_text


FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "serial_no": (
        "Serial No.",
        "Serial No",
        "Invoice No.",
        "Contract No.",
        "发票号",
        "合同号",
    ),
    "qad_pn": ("QAD PN", "Part No.", "Part No", "料号"),
    "quantity": ("Quantity", "Qty", "数量"),
    "amount": ("Amount", "金额"),
    "currency": ("Currency", "币种"),
    "pickup_date": (
        "Pick up Date",
        "Pickup Date",
        "提货日期",
        "日期",
    ),
    "ship_to": ("Ship to", "Ship To", "收货方", "收货人"),
    "delivery_way": (
        "Delivery Way",
        "Delivery Mode",
        "运输方式",
    ),
    "item_no": ("项号", "Item", "Item No."),
    "goods_name": (
        "商品名称及规格型号",
        "商品名称",
        "货物名称",
    ),
    "hs_code": ("商品编号", "HS Code", "HSCode", "海关编码"),
    "brand": ("品牌", "Brand"),
    "description_en": (
        "Description",
        "English Description",
        "英文描述",
    ),
    "unit_price": ("Unit Price", "UP", "单价"),
    "po_no": ("PO No.", "PO No", "PO_No", "PO"),
    "source_file": ("Source File", "来源文件", "源文件"),
}

ALIAS_MAP = {
    normalize_header(alias): field_name
    for field_name, aliases in FIELD_ALIASES.items()
    for alias in aliases
}


@dataclass(frozen=True)
class TemplateSpec:
    sheet_name: str
    header_row: int
    columns: dict[int, str]
    unknown_headers: tuple[str, ...]
    max_column: int


def inspect_template(path: str | Path) -> TemplateSpec:
    template_path = Path(path)
    workbook = load_workbook(template_path, data_only=False, read_only=True)
    candidates: list[
        tuple[int, int, int, int, str, dict[int, str], list[str], int]
    ] = []
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            for row_index, row in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=min(50, worksheet.max_row),
                    values_only=True,
                ),
                1,
            ):
                columns: dict[int, str] = {}
                unknown: list[str] = []
                for column_index, value in enumerate(row, 1):
                    header = safe_text(value)
                    if not header:
                        continue
                    field_name = ALIAS_MAP.get(normalize_header(header))
                    if field_name:
                        columns[column_index] = field_name
                    else:
                        unknown.append(header)
                recognized = len(columns)
                if "qad_pn" not in columns.values() or recognized < 2:
                    continue
                visible_bonus = 1 if worksheet.sheet_state == "visible" else 0
                candidates.append(
                    (
                        recognized,
                        visible_bonus,
                        -sheet_index,
                        -row_index,
                        worksheet.title,
                        columns,
                        unknown,
                        row_index,
                    )
                )
                break
    finally:
        workbook.close()

    if not candidates:
        raise ValueError(
            "输出模板中没有找到可识别的表头；至少需要 QAD PN/Part No. "
            "以及另一个支持字段"
        )
    candidates.sort(key=lambda item: item[:4], reverse=True)
    _, _, _, _, sheet_name, columns, unknown, header_row = candidates[0]

    workbook = load_workbook(template_path, data_only=False, read_only=True)
    try:
        max_column = workbook[sheet_name].max_column
    finally:
        workbook.close()

    return TemplateSpec(
        sheet_name=sheet_name,
        header_row=header_row,
        columns=columns,
        unknown_headers=tuple(unknown),
        max_column=max_column,
    )


def _copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def write_output(
    template_path: str | Path,
    spec: TemplateSpec,
    rows: list[dict[str, Any]],
    output_path: str | Path,
) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template_path, data_only=False)
    worksheet = workbook[spec.sheet_name]

    sample_row_index = spec.header_row + 1
    sample_cells = {
        column_index: copy(worksheet.cell(sample_row_index, column_index))
        for column_index in range(1, spec.max_column + 1)
    }
    if worksheet.max_row > spec.header_row:
        worksheet.delete_rows(
            spec.header_row + 1,
            worksheet.max_row - spec.header_row,
        )

    approximate_fill = PatternFill("solid", fgColor="FFF2CC")
    missing_fill = PatternFill("solid", fgColor="F4CCCC")
    rule_fields = {"goods_name", "hs_code", "brand"}

    for row_offset, values in enumerate(rows, 1):
        output_row = spec.header_row + row_offset
        for column_index, field_name in spec.columns.items():
            cell = worksheet.cell(output_row, column_index)
            sample = sample_cells.get(column_index)
            if sample is not None:
                _copy_cell_style(sample, cell)
            value = values.get(field_name, "")
            cell.value = value
            if field_name in {"qad_pn", "item_no", "hs_code", "serial_no"}:
                cell.number_format = "@"
            elif field_name == "pickup_date":
                cell.number_format = "yyyy-mm-dd"
            elif field_name == "amount":
                cell.number_format = "#,##0.00"
            elif field_name in {"quantity", "unit_price"}:
                cell.number_format = "#,##0.######"
            if field_name in rule_fields:
                match_type = values.get("_rule_match_type")
                if match_type == "naming_rule":
                    cell.fill = approximate_fill
                elif match_type in {"ambiguous", "missing"}:
                    cell.fill = missing_fill

    last_row = spec.header_row + max(1, len(rows))
    if worksheet.auto_filter.ref:
        worksheet.auto_filter.ref = (
            f"A{spec.header_row}:"
            f"{get_column_letter(spec.max_column)}{last_row}"
        )
    for table in worksheet.tables.values():
        table.ref = (
            f"A{spec.header_row}:"
            f"{get_column_letter(spec.max_column)}{last_row}"
        )
    if worksheet.freeze_panes is None:
        worksheet.freeze_panes = f"A{spec.header_row + 1}"

    workbook.save(output)
    workbook.close()

    verification = load_workbook(output, data_only=False, read_only=True)
    verification.close()
    return output

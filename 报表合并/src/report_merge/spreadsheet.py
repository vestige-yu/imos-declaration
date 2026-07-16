from __future__ import annotations

import struct
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .text_utils import safe_text


@dataclass(frozen=True)
class SheetData:
    name: str
    state: str
    rows: list[list[Any]]


class XlsWorkbook:
    """读取本业务使用的 BIFF8 老式 .xls 文件。"""

    END_OF_CHAIN = 0xFFFFFFFE
    FREE_SECTOR = 0xFFFFFFFF

    def __init__(self, data: bytes):
        self.data = data
        self.workbook_stream = self._read_workbook_stream()
        self.shared_strings = self._read_sst()
        self.sheet_metadata = self._read_sheet_metadata()
        self.sheets = self._read_sheets()

    @classmethod
    def load(cls, path: str | Path) -> "XlsWorkbook":
        return cls(Path(path).read_bytes())

    def _u16(self, offset: int) -> int:
        return struct.unpack_from("<H", self.data, offset)[0]

    def _u32(self, offset: int) -> int:
        return struct.unpack_from("<I", self.data, offset)[0]

    def _sector(self, index: int, sector_size: int) -> bytes:
        start = 512 + index * sector_size
        return self.data[start : start + sector_size]

    def _read_workbook_stream(self) -> bytes:
        if self.data[:8] != bytes.fromhex("d0cf11e0a1b11ae1"):
            raise ValueError("不是有效的老式 .xls 文件")
        sector_size = 1 << self._u16(30)
        fat_sector_count = self._u32(44)
        first_dir_sector = self._u32(48)
        difat = []
        for index in range(109):
            value = self._u32(76 + index * 4)
            if value != self.FREE_SECTOR:
                difat.append(value)

        fat: list[int] = []
        for sector_index in difat[:fat_sector_count]:
            block = self._sector(sector_index, sector_size)
            fat.extend(
                struct.unpack_from("<I", block, offset)[0]
                for offset in range(0, len(block), 4)
            )

        def chain(start: int) -> list[int]:
            result: list[int] = []
            seen: set[int] = set()
            current = start
            while (
                current not in (self.END_OF_CHAIN, self.FREE_SECTOR)
                and current not in seen
                and current < len(fat)
            ):
                seen.add(current)
                result.append(current)
                current = fat[current]
            return result

        directory = b"".join(
            self._sector(index, sector_size) for index in chain(first_dir_sector)
        )
        workbook_start = None
        workbook_size = None
        for offset in range(0, len(directory), 128):
            entry = directory[offset : offset + 128]
            if len(entry) < 128:
                continue
            name_len = struct.unpack_from("<H", entry, 64)[0]
            name = entry[: max(0, name_len - 2)].decode("utf-16le", "ignore")
            if name in ("Workbook", "Book"):
                workbook_start = struct.unpack_from("<I", entry, 116)[0]
                workbook_size = struct.unpack_from("<I", entry, 120)[0]
                break
        if workbook_start is None or workbook_size is None:
            raise ValueError("未找到 Workbook 数据流")
        stream = b"".join(
            self._sector(index, sector_size) for index in chain(workbook_start)
        )
        return stream[:workbook_size]

    def _records(self):
        position = 0
        while position + 4 <= len(self.workbook_stream):
            record_type, size = struct.unpack_from(
                "<HH", self.workbook_stream, position
            )
            payload = self.workbook_stream[position + 4 : position + 4 + size]
            yield record_type, payload
            position += 4 + size

    def _read_sst(self) -> list[str]:
        segments: list[bytes] = []
        collecting = False
        for record_type, payload in self._records():
            if record_type == 0x00FC:
                segments = [payload]
                collecting = True
            elif record_type == 0x003C and collecting:
                segments.append(payload)
            elif collecting:
                break
        if not segments or len(segments[0]) < 8:
            return []

        class SegmentReader:
            def __init__(self, source: list[bytes]):
                self.source = source
                self.segment = 0
                self.position = 8

            def _advance(self) -> bool:
                self.segment += 1
                self.position = 0
                return self.segment < len(self.source)

            def read(self, count: int) -> bytes:
                result = bytearray()
                while count > 0 and self.segment < len(self.source):
                    current = self.source[self.segment]
                    if self.position >= len(current):
                        if not self._advance():
                            break
                        continue
                    take = min(count, len(current) - self.position)
                    result.extend(current[self.position : self.position + take])
                    self.position += take
                    count -= take
                return bytes(result)

            def read_characters(self, count: int, is_16bit: bool) -> str:
                parts: list[str] = []
                remaining = count
                while remaining > 0 and self.segment < len(self.source):
                    current = self.source[self.segment]
                    if self.position >= len(current):
                        if not self._advance():
                            break
                        if self.position < len(self.source[self.segment]):
                            flags = self.source[self.segment][self.position]
                            self.position += 1
                            is_16bit = bool(flags & 0x01)
                        continue
                    width = 2 if is_16bit else 1
                    chars_here = min(
                        remaining, (len(current) - self.position) // width
                    )
                    if chars_here <= 0:
                        self.position = len(current)
                        continue
                    raw = current[
                        self.position : self.position + chars_here * width
                    ]
                    self.position += chars_here * width
                    remaining -= chars_here
                    parts.append(
                        raw.decode("utf-16le" if is_16bit else "latin1", "ignore")
                    )
                return "".join(parts)

        unique_count = struct.unpack_from("<I", segments[0], 4)[0]
        reader = SegmentReader(segments)
        strings: list[str] = []
        for _ in range(unique_count):
            header = reader.read(3)
            if len(header) < 3:
                break
            char_count = struct.unpack_from("<H", header, 0)[0]
            flags = header[2]
            rich_runs = 0
            ext_size = 0
            if flags & 0x08:
                raw = reader.read(2)
                if len(raw) == 2:
                    rich_runs = struct.unpack("<H", raw)[0]
            if flags & 0x04:
                raw = reader.read(4)
                if len(raw) == 4:
                    ext_size = struct.unpack("<I", raw)[0]
            strings.append(
                reader.read_characters(char_count, bool(flags & 0x01))
            )
            if rich_runs:
                reader.read(rich_runs * 4)
            if ext_size:
                reader.read(ext_size)
        return strings

    @staticmethod
    def _decode_rk(raw: int) -> float:
        multiplied = raw & 0x01
        is_integer = raw & 0x02
        value_bits = raw & 0xFFFFFFFC
        if is_integer:
            if value_bits & 0x80000000:
                value_bits -= 0x100000000
            value = value_bits >> 2
        else:
            value = struct.unpack("<d", struct.pack("<Q", value_bits << 32))[0]
        return value / 100 if multiplied else value

    @staticmethod
    def _formula_result(raw: bytes) -> float | None:
        if len(raw) < 8 or raw[6:8] == b"\xFF\xFF":
            return None
        value = struct.unpack("<d", raw[:8])[0]
        return None if value != value else value

    @staticmethod
    def _decode_label_text(payload: bytes) -> str:
        if len(payload) < 8:
            return ""
        char_count = struct.unpack_from("<H", payload, 6)[0]
        remaining = len(payload) - 8
        if remaining >= char_count and (
            remaining == char_count
            or payload[8] not in (0, 1, 4, 5, 8, 9, 12, 13)
        ):
            return payload[8 : 8 + char_count].decode("latin1", "ignore")
        if remaining >= char_count * 2 and remaining == char_count * 2:
            return payload[8 : 8 + char_count * 2].decode(
                "utf-16le", "ignore"
            )
        if len(payload) < 9:
            return ""
        flags = payload[8]
        offset = 9
        width = 2 if flags & 0x01 else 1
        raw = payload[offset : offset + char_count * width]
        return raw.decode("utf-16le" if flags & 0x01 else "latin1", "ignore")

    def _read_sheet_metadata(self) -> list[dict[str, str]]:
        sheets: list[dict[str, str]] = []
        for record_type, payload in self._records():
            if record_type != 0x0085 or len(payload) < 8:
                continue
            name_length = payload[6]
            flags = payload[7]
            width = 2 if flags & 0x01 else 1
            raw = payload[8 : 8 + name_length * width]
            sheets.append(
                {
                    "name": raw.decode(
                        "utf-16le" if width == 2 else "latin1", "ignore"
                    ),
                    "state": "visible" if payload[4] == 0 else "hidden",
                }
            )
        return sheets

    def _read_sheets(self) -> list[dict[tuple[int, int], Any]]:
        sheets: list[dict[tuple[int, int], Any]] = []
        current: dict[tuple[int, int], Any] | None = None
        pending_formula_cell: tuple[int, int] | None = None
        for record_type, payload in self._records():
            if record_type == 0x0809 and len(payload) >= 4:
                stream_type = struct.unpack_from("<H", payload, 2)[0]
                if stream_type == 0x0010:
                    current = {}
                    pending_formula_cell = None
            elif record_type == 0x000A:
                if current is not None:
                    sheets.append(current)
                    current = None
                pending_formula_cell = None
            elif current is None:
                continue
            elif record_type == 0x0204 and len(payload) >= 9:
                row, column, _ = struct.unpack_from("<HHH", payload, 0)
                current[(row, column)] = self._decode_label_text(payload)
            elif record_type == 0x00FD and len(payload) >= 10:
                row, column, _, index = struct.unpack_from("<HHHI", payload, 0)
                current[(row, column)] = (
                    self.shared_strings[index]
                    if index < len(self.shared_strings)
                    else ""
                )
            elif record_type == 0x0203 and len(payload) >= 14:
                row, column, _ = struct.unpack_from("<HHH", payload, 0)
                current[(row, column)] = struct.unpack_from(
                    "<d", payload, 6
                )[0]
            elif record_type == 0x027E and len(payload) >= 10:
                row, column, _ = struct.unpack_from("<HHH", payload, 0)
                current[(row, column)] = self._decode_rk(
                    struct.unpack_from("<I", payload, 6)[0]
                )
            elif record_type == 0x00BD and len(payload) >= 6:
                row, first_column, last_column = struct.unpack_from(
                    "<HHH", payload, 0
                )
                offset = 6
                for column in range(first_column, last_column + 1):
                    if offset + 6 <= len(payload):
                        _, raw = struct.unpack_from("<HI", payload, offset)
                        current[(row, column)] = self._decode_rk(raw)
                    offset += 6
            elif record_type == 0x0006 and len(payload) >= 14:
                row, column, _ = struct.unpack_from("<HHH", payload, 0)
                value = self._formula_result(payload[6:14])
                pending_formula_cell = (row, column)
                if value is not None:
                    current[(row, column)] = value
            elif (
                record_type == 0x0207
                and pending_formula_cell
                and len(payload) >= 3
            ):
                length = struct.unpack_from("<H", payload, 0)[0]
                flags = payload[2]
                raw = payload[3 : 3 + length * (2 if flags & 1 else 1)]
                current[pending_formula_cell] = raw.decode(
                    "utf-16le" if flags & 1 else "latin1", "ignore"
                )
                pending_formula_cell = None
        return sheets


def matrix_from_cells(cells: dict[tuple[int, int], Any]) -> list[list[Any]]:
    if not cells:
        return []
    max_row = max(row for row, _ in cells)
    max_column = max(column for _, column in cells)
    return [
        [cells.get((row, column), "") for column in range(max_column + 1)]
        for row in range(max_row + 1)
    ]


def read_workbook(path: str | Path) -> list[SheetData]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        workbook = XlsWorkbook.load(path)
        result: list[SheetData] = []
        for index, cells in enumerate(workbook.sheets):
            metadata = (
                workbook.sheet_metadata[index]
                if index < len(workbook.sheet_metadata)
                else {"name": f"Sheet {index + 1}", "state": "visible"}
            )
            result.append(
                SheetData(
                    name=metadata["name"],
                    state=metadata["state"],
                    rows=matrix_from_cells(cells),
                )
            )
        return result
    if suffix == ".xlsx":
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            result = []
            for worksheet in workbook.worksheets:
                rows = [
                    list(row)
                    for row in worksheet.iter_rows(values_only=True)
                ]
                result.append(
                    SheetData(
                        name=worksheet.title,
                        state=worksheet.sheet_state,
                        rows=rows,
                    )
                )
            return result
        finally:
            workbook.close()
    raise ValueError(f"不支持的文件格式：{path.name}；仅支持 .xls 和 .xlsx")

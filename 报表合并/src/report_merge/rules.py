from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .models import MergeWarning, RuleRecord
from .text_utils import (
    canonical_part_number,
    find_column,
    normalize_hs_code,
    normalize_part,
    part_match_rule_labels,
    safe_text,
)


@dataclass(frozen=True)
class RuleMatch:
    record: RuleRecord | None
    match_type: str
    applied_rules: tuple[str, ...] = ()
    candidates: tuple[RuleRecord, ...] = ()


@dataclass(frozen=True)
class RuleSet:
    sheet_name: str
    records: dict[str, RuleRecord]
    canonical_records: dict[str, RuleRecord]
    ambiguous_records: dict[str, tuple[RuleRecord, ...]]
    warnings: tuple[MergeWarning, ...]

    def match(
        self,
        qad_pn: str,
        allow_naming_rule_match: bool = True,
    ) -> RuleMatch:
        key = normalize_part(qad_pn)
        exact = self.records.get(key)
        if exact is not None:
            return RuleMatch(exact, "exact")
        if not allow_naming_rule_match:
            return RuleMatch(None, "missing")
        canonical_key = canonical_part_number(qad_pn)
        ambiguous = self.ambiguous_records.get(canonical_key)
        if ambiguous is not None:
            return RuleMatch(
                None,
                "ambiguous",
                candidates=ambiguous,
            )
        matched = self.canonical_records.get(canonical_key)
        if matched is not None:
            return RuleMatch(
                matched,
                "naming_rule",
                applied_rules=part_match_rule_labels(
                    qad_pn,
                    matched.qad_pn,
                ),
            )
        return RuleMatch(None, "missing")


def _candidate_score(
    headers: list[object],
    records: list[RuleRecord],
) -> float:
    part_column = find_column(
        headers,
        ("QAD PN", "Part No.", "Part No", "PartNo", "料号"),
    )
    goods_column = find_column(
        headers,
        ("商品名称及规格型号", "商品名称", "Description", "货物名称"),
    )
    hs_column = find_column(
        headers,
        ("商品编号", "HS Code", "HSCode", "海关编码"),
    )
    score = 0.0
    if part_column is not None:
        score += 30
    if goods_column is not None:
        score += 20
    if hs_column is not None:
        score += 20
    if records:
        unique_count = len({record.qad_pn for record in records})
        score += 20 * unique_count / len(records)
        score += min(10, unique_count / 50)
    return score


def load_rules(path: str | Path) -> RuleSet:
    rule_path = Path(path)
    workbook = load_workbook(rule_path, data_only=True, read_only=True)
    candidates: list[
        tuple[float, str, list[RuleRecord], list[MergeWarning]]
    ] = []
    try:
        for worksheet in workbook.worksheets:
            raw_rows = list(worksheet.iter_rows(values_only=True))
            for header_index, raw_headers in enumerate(raw_rows[:50]):
                headers = list(raw_headers)
                part_column = find_column(
                    headers,
                    ("QAD PN", "Part No.", "Part No", "PartNo", "料号"),
                )
                goods_column = find_column(
                    headers,
                    (
                        "商品名称及规格型号",
                        "商品名称",
                        "Description",
                        "货物名称",
                    ),
                )
                hs_column = find_column(
                    headers,
                    ("商品编号", "HS Code", "HSCode", "海关编码"),
                )
                if part_column is None or (
                    goods_column is None and hs_column is None
                ):
                    continue
                brand_column = find_column(headers, ("品牌", "Brand"))
                note_column = find_column(
                    headers,
                    ("备注", "Note", "状态", "Status"),
                )
                records: list[RuleRecord] = []
                for source_row, raw_row in enumerate(
                    raw_rows[header_index + 1 :],
                    header_index + 2,
                ):
                    row = list(raw_row)
                    qad_pn = normalize_part(
                        row[part_column]
                        if part_column < len(row)
                        else ""
                    )
                    if not qad_pn:
                        continue
                    records.append(
                        RuleRecord(
                            qad_pn=qad_pn,
                            goods_name=safe_text(
                                row[goods_column]
                                if goods_column is not None
                                and goods_column < len(row)
                                else ""
                            ),
                            hs_code=normalize_hs_code(
                                row[hs_column]
                                if hs_column is not None
                                and hs_column < len(row)
                                else ""
                            ),
                            brand=safe_text(
                                row[brand_column]
                                if brand_column is not None
                                and brand_column < len(row)
                                else ""
                            ),
                            note=safe_text(
                                row[note_column]
                                if note_column is not None
                                and note_column < len(row)
                                else ""
                            ),
                            source_sheet=worksheet.title,
                            source_row=source_row,
                        )
                    )
                if not records:
                    continue
                candidates.append(
                    (
                        _candidate_score(headers, records),
                        worksheet.title,
                        records,
                        [],
                    )
                )
                break
    finally:
        workbook.close()

    if not candidates:
        raise ValueError(
            "规则文件中没有找到 QAD PN/Part No. 与商品名称或商品编号的映射表"
        )

    candidates.sort(key=lambda item: item[0], reverse=True)
    _, sheet_name, selected_records, warnings = candidates[0]
    record_map: dict[str, RuleRecord] = {}
    for record in selected_records:
        existing = record_map.get(record.qad_pn)
        if existing is not None and (
            existing.goods_name != record.goods_name
            or existing.hs_code != record.hs_code
        ):
            warnings.append(
                MergeWarning(
                    code="RULE_DUPLICATE_CONFLICT",
                    message=(
                        f"规则表 {sheet_name} 中 {record.qad_pn} 存在冲突；"
                        f"已采用第 {record.source_row} 行"
                    ),
                    qad_pn=record.qad_pn,
                )
            )
        record_map[record.qad_pn] = record

    canonical_groups: dict[str, list[RuleRecord]] = {}
    for record in record_map.values():
        canonical_groups.setdefault(
            canonical_part_number(record.qad_pn),
            [],
        ).append(record)

    canonical_records: dict[str, RuleRecord] = {}
    ambiguous_records: dict[str, tuple[RuleRecord, ...]] = {}
    for canonical_key, grouped_records in canonical_groups.items():
        mapping_signatures = {
            (
                record.goods_name,
                record.hs_code,
                record.brand,
            )
            for record in grouped_records
        }
        if len(mapping_signatures) == 1:
            canonical_records[canonical_key] = grouped_records[0]
        else:
            ambiguous_records[canonical_key] = tuple(grouped_records)

    return RuleSet(
        sheet_name=sheet_name,
        records=record_map,
        canonical_records=canonical_records,
        ambiguous_records=ambiguous_records,
        warnings=tuple(warnings),
    )

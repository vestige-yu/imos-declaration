from __future__ import annotations

import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

import app  # noqa: E402


def build_source(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoice"
    worksheet["H4"] = "Serial No:"
    worksheet["I4"] = "SPSELFTEST"
    worksheet["H5"] = "pick up Date:"
    worksheet["I5"] = "2026-01-16"
    worksheet["G16"] = "Ship to:"
    worksheet["G17"] = "SELF TEST CUSTOMER"
    worksheet["G24"] = "Delivery way:Sea"
    headers = (
        "Item",
        "QAD PN",
        "Description",
        "Quantity",
        "UP",
        "PO_No",
        "Amount",
        "Currency",
    )
    for column, value in enumerate(headers, 2):
        worksheet.cell(26, column, value)
    values = (
        "1",
        "2010300096-00",
        "SELF TEST PART",
        3,
        4.113,
        "PO-SELFTEST",
        12.34,
        "USD",
    )
    for column, value in enumerate(values, 2):
        worksheet.cell(27, column, value)
    workbook.save(path)
    workbook.close()


class DesktopAppTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        root = Path(self.temporary.name)
        self.original_paths = {
            name: getattr(app, name)
            for name in (
                "DATA_DIR",
                "TEMPLATE_DIR",
                "RULES_DIR",
                "HISTORY_DIR",
                "OUTPUT_DIR",
                "TEMP_DIR",
                "DB_PATH",
                "META_PATH",
            )
        }
        app.DATA_DIR = root
        app.TEMPLATE_DIR = root / "templates"
        app.RULES_DIR = root / "rules"
        app.HISTORY_DIR = root / "history"
        app.OUTPUT_DIR = root / "outputs"
        app.TEMP_DIR = root / "temp"
        app.DB_PATH = root / "app.db"
        app.META_PATH = root / "config.json"
        app.SESSIONS.clear()
        self.source_path = root / "source.xlsx"
        build_source(self.source_path)

    def tearDown(self) -> None:
        for name, path in self.original_paths.items():
            setattr(app, name, path)
        app.SESSIONS.clear()
        self.temporary.cleanup()

    def test_preview_generate_and_history_archive(self) -> None:
        record_id, preview = app.create_preview_record([self.source_path])
        summary = preview["summary"]
        self.assertEqual(summary["sourceCount"], 1)
        self.assertEqual(summary["rowCount"], 1)
        self.assertEqual(summary["exactCount"], 1)
        self.assertEqual(summary["namingRuleCount"], 0)
        self.assertEqual(summary["missingCount"], 0)
        self.assertFalse(summary["requiresConfirmation"])

        output_path, output_name, generated = (
            app.generate_history_record(record_id)
        )
        self.assertTrue(output_path.is_file())
        self.assertEqual(generated["rowCount"], 1)
        self.assertTrue(output_name.endswith(".xlsx"))

        row = app._history_row(record_id)
        history = app.row_to_history(row, include_preview=True)
        self.assertEqual(history["status"], "generated")
        self.assertEqual(history["sourceCount"], 1)
        self.assertEqual(history["warningCount"], 0)
        record_dir = app.HISTORY_DIR / record_id
        self.assertTrue((record_dir / "template.xlsx").is_file())
        self.assertTrue((record_dir / "rules.xlsx").is_file())
        self.assertTrue((record_dir / "preview.json").is_file())
        self.assertTrue((record_dir / "output.xlsx").is_file())
        self.assertEqual(
            len(list((record_dir / "sources").glob("*/*"))),
            1,
        )

    def test_uploaded_config_is_validated_versioned_and_activated(self) -> None:
        temporary_template = Path(self.temporary.name) / "template.xlsx"
        temporary_rules = Path(self.temporary.name) / "rules.xlsx"
        shutil.copy2(app.DEFAULT_TEMPLATE, temporary_template)
        shutil.copy2(app.DEFAULT_RULES, temporary_rules)
        result = app.update_active_config(
            (temporary_template, "客户模板.xlsx"),
            (temporary_rules, "客户规则.xlsx"),
        )
        self.assertEqual(
            result["active"]["template"]["filename"],
            "客户模板.xlsx",
        )
        self.assertEqual(
            result["active"]["rules"]["filename"],
            "客户规则.xlsx",
        )
        self.assertTrue((app.TEMPLATE_DIR / "current.xlsx").is_file())
        self.assertTrue((app.RULES_DIR / "current.xlsx").is_file())
        self.assertEqual(
            len(list((app.TEMPLATE_DIR / "versions").glob("*.xlsx"))),
            1,
        )
        self.assertEqual(
            len(list((app.RULES_DIR / "versions").glob("*.xlsx"))),
            1,
        )


if __name__ == "__main__":
    unittest.main()

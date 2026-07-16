from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from report_merge import MergeConfig, merge_reports  # noqa: E402
from report_merge.rules import load_rules  # noqa: E402
from report_merge.source_parser import parse_source_report  # noqa: E402
from report_merge.text_utils import canonical_part_number  # noqa: E402


SOURCE_FILES = sorted(PROJECT_ROOT.glob("测试表*.xls"))
TEMPLATE = PROJECT_ROOT / "2026 Daily Export List模板.xlsx"
RULES = PROJECT_ROOT / "List模板.xlsx"


class SourceParserTests(unittest.TestCase):
    @unittest.skipUnless(
        len(SOURCE_FILES) == 3,
        "本地真实合并样例未包含在当前检出目录中",
    )
    def test_actual_sources_have_expected_totals(self) -> None:
        expected = {
            "SP26000096": (24, 112546.0, 246950.57),
            "SP26002566": (4, 7470.0, 40791.64),
            "SP26003466": (9, 5415.0, 61271.58),
        }
        self.assertEqual(len(SOURCE_FILES), 3)
        for source_file in SOURCE_FILES:
            report = parse_source_report(source_file)
            rows, quantity, amount = expected[report.serial_no]
            self.assertEqual(len(report.lines), rows)
            self.assertAlmostEqual(
                sum(line.quantity for line in report.lines),
                quantity,
            )
            self.assertAlmostEqual(
                sum(line.amount for line in report.lines),
                amount,
                places=2,
            )
            self.assertTrue(report.ship_to)
            self.assertEqual(report.delivery_way, "Sea")


class RuleTests(unittest.TestCase):
    def test_confirmed_part_number_naming_rules(self) -> None:
        same_pairs = (
            ("3073400200-00", "307340020000"),
            ("001234", "1234"),
            ("123400", "1234"),
            ("ABC001-01", "ABC001-02"),
            ("ＡＢＣ　００１", "ABC001"),
            ("123456.0", "123456"),
        )
        different_pairs = (
            ("CABC001", "ABC001"),
            ("ABO001", "AB0001"),
            ("ABI001", "AB1001"),
            ("OLD123", "NEW456"),
        )
        for left, right in same_pairs:
            self.assertEqual(
                canonical_part_number(left),
                canonical_part_number(right),
                (left, right),
            )
        for left, right in different_pairs:
            self.assertNotEqual(
                canonical_part_number(left),
                canonical_part_number(right),
                (left, right),
            )

    def test_rules_choose_sheet2_and_apply_confirmed_naming_rules(self) -> None:
        rules = load_rules(RULES)
        self.assertEqual(rules.sheet_name, "Sheet2")
        self.assertEqual(len(rules.ambiguous_records), 6)

        exact = rules.match("2010300096-00")
        self.assertIsNotNone(exact.record)
        self.assertEqual(exact.match_type, "exact")

        naming_match = rules.match("3031100045-01A")
        self.assertIsNotNone(naming_match.record)
        self.assertEqual(naming_match.record.qad_pn, "3031100045-01")
        self.assertEqual(naming_match.match_type, "naming_rule")
        self.assertIn("忽略版本后缀差异", naming_match.applied_rules)

        exact_conflicting_variant = rules.match("3020200324-01C")
        self.assertEqual(exact_conflicting_variant.match_type, "exact")
        self.assertEqual(
            exact_conflicting_variant.record.hs_code,
            "87082990",
        )

        ambiguous = rules.match("3020200324-01B")
        self.assertIsNone(ambiguous.record)
        self.assertEqual(ambiguous.match_type, "ambiguous")
        self.assertEqual(len(ambiguous.candidates), 3)

        missing = rules.match("NOT-EXIST")
        self.assertIsNone(missing.record)
        self.assertEqual(missing.match_type, "missing")


class MergeEngineTests(unittest.TestCase):
    @staticmethod
    def _build_source(path: Path, qad_pn: str) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["H4"] = "Serial No:"
        worksheet["I4"] = "SPTEST"
        worksheet["H5"] = "pick up Date:"
        worksheet["I5"] = "2026-01-16"
        worksheet["G16"] = "Ship to:"
        worksheet["G17"] = "TEST CUSTOMER"
        worksheet["G24"] = "Delivery way:Sea"
        headers = (
            "Item",
            "QAD PN",
            "Description",
            "Quantity",
            "UP",
            "PO_No",
            "Amount",
            "Currency",
        )
        for column, value in enumerate(headers, 2):
            worksheet.cell(26, column, value)
        values = (
            "1",
            qad_pn,
            "TEST PART",
            1,
            1,
            "PO-TEST",
            1,
            "USD",
        )
        for column, value in enumerate(values, 2):
            worksheet.cell(27, column, value)
        workbook.save(path)
        workbook.close()

    @unittest.skipUnless(
        len(SOURCE_FILES) == 3,
        "本地真实合并样例未包含在当前检出目录中",
    )
    def test_actual_files_merge_to_expected_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "merged.xlsx"
            result = merge_reports(
                SOURCE_FILES,
                TEMPLATE,
                RULES,
                output,
            )
            self.assertEqual(result.row_count, 37)
            self.assertAlmostEqual(result.total_quantity, 125431.0)
            self.assertAlmostEqual(result.total_amount, 349013.79, places=2)
            self.assertEqual(result.rule_sheet, "Sheet2")
            self.assertEqual(
                [warning.code for warning in result.warnings],
                ["RULE_NAMING_MATCH"],
            )

            workbook = load_workbook(output, data_only=False)
            worksheet = workbook["Sheet1"]
            self.assertEqual(worksheet.max_row, 38)
            self.assertEqual(worksheet["A2"].value, "SP26000096")
            self.assertEqual(worksheet["B7"].value, "3031100045-01A")
            self.assertEqual(
                worksheet["G2"].value,
                "CASCO Imos Italia S.R.L.",
            )
            self.assertEqual(worksheet["G26"].value, "CASCO Logistics GmbH")
            self.assertEqual(worksheet["K2"].value, "8708299000")
            self.assertEqual(worksheet["K2"].number_format, "@")
            self.assertEqual(worksheet["J7"].fill.fill_type, "solid")
            self.assertEqual(worksheet["J7"].fill.fgColor.rgb, "00FFF2CC")
            workbook.close()

    def test_more_than_30_files_is_rejected_before_parsing(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            self._build_source(source, "2010300096-00")
            with self.assertRaisesRegex(ValueError, "最多处理 30 份"):
                merge_reports(
                    [source] * 31,
                    TEMPLATE,
                    RULES,
                    Path(temp_dir) / "merged.xlsx",
                )

    def test_ambiguous_naming_group_is_not_selected_automatically(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "ambiguous.xlsx"
            output = Path(temp_dir) / "merged.xlsx"
            self._build_source(source, "3020200324-01B")
            result = merge_reports(
                [source],
                TEMPLATE,
                RULES,
                output,
            )
            self.assertEqual(
                [warning.code for warning in result.warnings],
                ["RULE_NAMING_AMBIGUOUS"],
            )
            self.assertEqual(
                result.rows[0]["_rule_match_type"],
                "ambiguous",
            )
            self.assertEqual(result.rows[0]["goods_name"], "")
            self.assertEqual(result.rows[0]["hs_code"], "")

            workbook = load_workbook(output, data_only=False)
            worksheet = workbook["Sheet1"]
            self.assertIsNone(worksheet["J2"].value)
            self.assertIsNone(worksheet["K2"].value)
            self.assertEqual(worksheet["J2"].fill.fill_type, "solid")
            self.assertEqual(worksheet["J2"].fill.fgColor.rgb, "00F4CCCC")
            workbook.close()

    def test_block_policy_stops_on_unmatched_rule(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "missing.xlsx"
            self._build_source(source, "9999999999")
            with self.assertRaisesRegex(ValueError, "已按 block 策略停止生成"):
                merge_reports(
                    [source],
                    TEMPLATE,
                    RULES,
                    Path(temp_dir) / "merged.xlsx",
                    MergeConfig(
                        unmatched_policy="block",
                        allow_naming_rule_match=False,
                    ),
                )


if __name__ == "__main__":
    unittest.main()

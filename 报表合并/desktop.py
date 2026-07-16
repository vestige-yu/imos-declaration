#!/usr/bin/env python3
from __future__ import annotations

import os
import socket
import sys
import tempfile
import threading
import traceback
import webbrowser
from pathlib import Path

from openpyxl import Workbook, load_workbook

import app
from report_merge import merge_reports


def log_startup_error(message: str) -> None:
    try:
        app.ensure_app_dirs()
        with (app.DATA_DIR / "startup.log").open(
            "a",
            encoding="utf-8",
        ) as handle:
            handle.write(message.rstrip() + "\n\n")
    except Exception:
        pass


def find_free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def open_browser_and_wait(url: str) -> None:
    webbrowser.open(url)
    threading.Event().wait()


def _build_self_test_source(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoice"
    worksheet["H4"] = "Serial No:"
    worksheet["I4"] = "SPSELFTEST"
    worksheet["H5"] = "pick up Date:"
    worksheet["I5"] = "2026-01-16"
    worksheet["G16"] = "Ship to:"
    worksheet["G17"] = "SELF TEST CUSTOMER"
    worksheet["G24"] = "Delivery way:Sea"
    headers = [
        "",
        "Item",
        "QAD PN",
        "Customer P/N",
        "Description",
        "Quantity",
        "UP",
        "PO_No",
        "Amount",
        "Currency",
    ]
    for column, value in enumerate(headers, 1):
        worksheet.cell(26, column, value)
    values = [
        "",
        "1",
        "2010300096-00",
        "2010300096-00",
        "SELF TEST PART",
        3,
        4.113,
        "PO-SELFTEST",
        12.34,
        "USD",
    ]
    for column, value in enumerate(values, 1):
        worksheet.cell(27, column, value)
    workbook.save(path)
    workbook.close()


def self_test() -> int:
    try:
        app.config_status()
        with tempfile.TemporaryDirectory() as temporary:
            temp_dir = Path(temporary)
            source = temp_dir / "self-test.xlsx"
            output = temp_dir / "self-test-output.xlsx"
            _build_self_test_source(source)
            result = merge_reports(
                [source],
                app.active_template_path(),
                app.active_rules_path(),
                output,
            )
            if result.row_count != 1:
                raise RuntimeError("自检输出明细数不正确")
            workbook = load_workbook(output, data_only=False, read_only=True)
            worksheet = workbook[result.template_sheet]
            if worksheet.cell(2, 2).value != "2010300096-00":
                raise RuntimeError("自检输出料号不正确")
            workbook.close()
        print("SELF_TEST_OK")
        return 0
    except Exception:
        log_startup_error("self-test failed:\n" + traceback.format_exc())
        traceback.print_exc()
        return 1


def main() -> int:
    if "--self-test" in sys.argv:
        return self_test()

    port = int(os.environ.get("PORT") or find_free_port())
    try:
        server, url = app.run_in_thread(
            host="127.0.0.1",
            port=port,
        )
    except Exception:
        log_startup_error(
            "server startup failed:\n" + traceback.format_exc()
        )
        raise

    try:
        log_startup_error(f"started successfully: {url}")
        open_browser_and_wait(url)
    finally:
        server.shutdown()
        server.server_close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

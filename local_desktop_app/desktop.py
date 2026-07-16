#!/usr/bin/env python3
import os
import socket
import sys
import tempfile
import threading
import traceback
import webbrowser
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook

import app


def log_startup_error(message):
    try:
        app.DATA_DIR.mkdir(parents=True, exist_ok=True)
        log_path = app.DATA_DIR / "startup.log"
        with log_path.open("a", encoding="utf-8") as handle:
            handle.write(message.rstrip() + "\n\n")
    except Exception:
        pass


def find_free_port():
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def open_browser_and_wait(url):
    webbrowser.open(url)
    threading.Event().wait()


def verify_xlsx_package(path):
    app.validate_xlsx_file(path, require_dimensions=True)
    with zipfile.ZipFile(path) as workbook:
        bad_file = workbook.testzip()
        if bad_file:
            raise RuntimeError(f"生成的 xlsx 压缩结构异常: {bad_file}")
        names = set(workbook.namelist())
        required = {
            "[Content_Types].xml",
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
            "xl/worksheets/sheet1.xml",
        }
        missing = required - names
        if missing:
            raise RuntimeError("生成的 xlsx 缺少必要文件: " + ", ".join(sorted(missing)))
        for name in workbook.namelist():
            if name.endswith(".xml"):
                ET.fromstring(workbook.read(name))


def declaration_self_test():
    preview = {
        "contractNo": "SELFTEST",
        "exportDateSerial": "",
        "consignee": "SELF TEST CUSTOMER",
        "packageKind": "PALLET",
        "packageCount": 1,
        "grossWeight": 2.5,
        "netWeight": 2.0,
        "tradeTerm": "FOB",
        "originCountry": "中国",
        "destinationCountry": "意大利",
        "domesticSource": "苏州",
        "totals": {"quantity": 3, "amount": 12.34, "currency": "USD"},
        "commodityLines": [
            {
                "itemNo": "1",
                "hsCode": "8708999990",
                "goodsName": "汽车零部件",
                "quantity": 3,
                "amount": 12.34,
                "currency": "USD",
                "brand": "示例品牌",
                "netWeight": 2.0,
            }
        ],
        "auditSamples": [
            {
                "itemNo": "1",
                "hsCode": "8708999990",
                "goodsName": "汽车零部件",
                "qadPartNo": "QAD-SELFTEST",
                "imosPartNo": "DECLARATION-SELFTEST",
                "invoiceQuantity": 3,
                "unitPrice": 4.113,
                "invoiceAmount": 12.34,
                "currency": "USD",
                "packingNetWeight": 2.0,
                "packingGrossWeight": 2.5,
                "poNo": "PO-SELFTEST",
                "invoiceSourceSheet": "Sheet 1",
                "invoiceSourceRow": 1,
            }
        ],
    }
    output_path, _ = app.generate_workbook(preview)
    verify_xlsx_package(output_path)
    workbook = load_workbook(output_path, data_only=False, read_only=True)
    if not workbook.sheetnames:
        raise RuntimeError("报关单自检工作簿没有工作表")
    workbook.close()
    return output_path


def build_merge_self_test_source(path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoice"
    worksheet["H4"] = "Serial No:"
    worksheet["I4"] = "SPSELFTEST"
    worksheet["H5"] = "pick up Date:"
    worksheet["I5"] = "2026-01-16"
    worksheet["G16"] = "Ship to:"
    worksheet["G17"] = "SELF TEST CUSTOMER"
    worksheet["G24"] = "Delivery way:Sea"
    headers = [
        "",
        "Item",
        "QAD PN",
        "Customer P/N",
        "Description",
        "Quantity",
        "UP",
        "PO_No",
        "Amount",
        "Currency",
    ]
    for column, value in enumerate(headers, 1):
        worksheet.cell(26, column, value)
    values = [
        "",
        "1",
        "2010300096-00",
        "2010300096-00",
        "SELF TEST PART",
        3,
        4.113,
        "PO-SELFTEST",
        12.34,
        "USD",
    ]
    for column, value in enumerate(values, 1):
        worksheet.cell(27, column, value)
    workbook.save(path)
    workbook.close()


def merge_self_test():
    service = app.merge_service
    service.config_status()
    with tempfile.TemporaryDirectory() as temporary:
        temp_dir = Path(temporary)
        source_path = temp_dir / "merge-source.xlsx"
        output_path = temp_dir / "merge-output.xlsx"
        build_merge_self_test_source(source_path)
        result = service.merge_reports(
            [source_path],
            service.active_template_path(),
            service.active_rules_path(),
            output_path,
        )
        if result.row_count != 1:
            raise RuntimeError("报表合并自检输出明细数不正确")
        verify_xlsx_package(output_path)
        workbook = load_workbook(output_path, data_only=False, read_only=True)
        worksheet = workbook[result.template_sheet]
        if worksheet.cell(2, 2).value != "2010300096-00":
            raise RuntimeError("报表合并自检输出料号不正确")
        workbook.close()
        persisted_output = app.merge_service.OUTPUT_DIR / "self-test-output.xlsx"
        persisted_output.parent.mkdir(parents=True, exist_ok=True)
        persisted_output.write_bytes(output_path.read_bytes())
        return persisted_output


def self_test():
    try:
        declaration_output = declaration_self_test()
        merge_output = merge_self_test()
        print(f"SELF_TEST_OK declaration={declaration_output} merge={merge_output}")
        return 0
    except Exception:
        log_startup_error("self-test failed:\n" + traceback.format_exc())
        traceback.print_exc()
        return 1


def main():
    if "--self-test" in sys.argv:
        return self_test()

    port = int(os.environ.get("PORT") or find_free_port())
    try:
        server, url = app.run_in_thread(host="127.0.0.1", port=port)
    except Exception:
        log_startup_error("server startup failed:\n" + traceback.format_exc())
        raise

    try:
        log_startup_error(f"started successfully: {url}")
        open_browser_and_wait(url)
    finally:
        server.shutdown()
        server.server_close()


if __name__ == "__main__":
    raise SystemExit(main())

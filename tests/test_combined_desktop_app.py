from __future__ import annotations

import json
import socket
import sys
import tempfile
import unittest
import urllib.error
import urllib.request
import uuid
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
DESKTOP_DIR = ROOT / "local_desktop_app"
sys.path.insert(0, str(DESKTOP_DIR))

import app  # noqa: E402


DECLARATION_INVOICE = (
    ROOT
    / "样例与测试资料/test/"
    "SP26000096 Invoice -CASCO Imos Italia -sea 20260116-1.xls"
)
DECLARATION_PACKING = (
    ROOT
    / "样例与测试资料/test/"
    "SP26000096 Packing list -CASCO Imos Italia-sea 20260116-1.xls"
)
def free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def multipart_body(files: list[tuple[str, Path]]) -> tuple[bytes, str]:
    boundary = f"----SuriWork{uuid.uuid4().hex}"
    chunks: list[bytes] = []
    for field_name, path in files:
        chunks.extend(
            [
                f"--{boundary}\r\n".encode(),
                (
                    f'Content-Disposition: form-data; name="{field_name}"; '
                    f'filename="{path.name}"\r\n'
                ).encode("utf-8"),
                b"Content-Type: application/octet-stream\r\n\r\n",
                path.read_bytes(),
                b"\r\n",
            ]
        )
    chunks.append(f"--{boundary}--\r\n".encode())
    return b"".join(chunks), boundary


def build_merge_source(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoice"
    worksheet["H4"] = "Serial No:"
    worksheet["I4"] = "SPSELFTEST"
    worksheet["H5"] = "pick up Date:"
    worksheet["I5"] = "2026-01-16"
    worksheet["G16"] = "Ship to:"
    worksheet["G17"] = "SELF TEST CUSTOMER"
    worksheet["G24"] = "Delivery way:Sea"
    headers = (
        "Item",
        "QAD PN",
        "Description",
        "Quantity",
        "UP",
        "PO_No",
        "Amount",
        "Currency",
    )
    for column, value in enumerate(headers, 2):
        worksheet.cell(26, column, value)
    values = (
        "1",
        "2010300096-00",
        "SELF TEST PART",
        3,
        4.113,
        "PO-SELFTEST",
        12.34,
        "USD",
    )
    for column, value in enumerate(values, 2):
        worksheet.cell(27, column, value)
    workbook.save(path)
    workbook.close()


class CombinedDesktopAppTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        root = Path(self.temporary.name)
        self.original_declaration = {
            name: getattr(app, name)
            for name in (
                "DATA_DIR",
                "TEMPLATE_DIR",
                "RULES_DIR",
                "HISTORY_DIR",
                "OUTPUT_DIR",
                "DB_PATH",
            )
        }
        app.DATA_DIR = root / "declaration"
        app.TEMPLATE_DIR = app.DATA_DIR / "templates"
        app.RULES_DIR = app.DATA_DIR / "rules"
        app.HISTORY_DIR = app.DATA_DIR / "history"
        app.OUTPUT_DIR = app.DATA_DIR / "outputs"
        app.DB_PATH = app.DATA_DIR / "app.db"
        app.SESSIONS.clear()

        merge = app.merge_service
        self.original_merge = {
            name: getattr(merge, name)
            for name in (
                "DATA_DIR",
                "TEMPLATE_DIR",
                "RULES_DIR",
                "HISTORY_DIR",
                "OUTPUT_DIR",
                "TEMP_DIR",
                "DB_PATH",
                "META_PATH",
            )
        }
        merge.DATA_DIR = root / "merge"
        merge.TEMPLATE_DIR = merge.DATA_DIR / "templates"
        merge.RULES_DIR = merge.DATA_DIR / "rules"
        merge.HISTORY_DIR = merge.DATA_DIR / "history"
        merge.OUTPUT_DIR = merge.DATA_DIR / "outputs"
        merge.TEMP_DIR = merge.DATA_DIR / "temp"
        merge.DB_PATH = merge.DATA_DIR / "app.db"
        merge.META_PATH = merge.DATA_DIR / "config.json"
        merge.SESSIONS.clear()
        self.merge_source = root / "merge-source.xlsx"
        build_merge_source(self.merge_source)

        self.server, self.base_url = app.run_in_thread(port=free_port())

    def tearDown(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        for name, value in self.original_declaration.items():
            setattr(app, name, value)
        for name, value in self.original_merge.items():
            setattr(app.merge_service, name, value)
        app.SESSIONS.clear()
        app.merge_service.SESSIONS.clear()
        self.temporary.cleanup()

    def request_json(
        self,
        path: str,
        *,
        method: str = "GET",
        payload: dict | None = None,
        files: list[tuple[str, Path]] | None = None,
    ) -> dict:
        headers: dict[str, str] = {}
        body = None
        if files is not None:
            body, boundary = multipart_body(files)
            headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
        elif payload is not None:
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            headers["Content-Type"] = "application/json"
        request = urllib.request.Request(
            self.base_url.rstrip("/") + path,
            data=body,
            headers=headers,
            method=method,
        )
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                result = json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as error:
            message = error.read().decode("utf-8", "replace")
            self.fail(f"{method} {path} 返回 {error.code}: {message}")
        self.assertTrue(result.get("ok"), result)
        return result

    def download(self, path: str) -> bytes:
        with urllib.request.urlopen(
            self.base_url.rstrip("/") + path,
            timeout=30,
        ) as response:
            return response.read()

    def assert_xlsx(self, content: bytes) -> None:
        path = Path(self.temporary.name) / f"{uuid.uuid4().hex}.xlsx"
        path.write_bytes(content)
        with zipfile.ZipFile(path) as archive:
            self.assertIsNone(archive.testzip())
            self.assertIn("xl/workbook.xml", archive.namelist())
        workbook = load_workbook(path, read_only=True, data_only=False)
        self.assertTrue(workbook.sheetnames)
        workbook.close()

    def test_frontend_and_four_config_files_are_isolated(self) -> None:
        page = self.download("/").decode("utf-8")
        self.assertIn("报关单生成", page)
        self.assertIn("报表合并", page)
        self.assertIn("合并后的生成模板", page)

        declaration = self.request_json(
            "/api/admin/rules",
            method="POST",
            files=[
                ("template", app.DEFAULT_TEMPLATE),
                ("rules", ROOT / "最新逻辑_请优先看/报关单配置关系表.xlsx"),
            ],
        )
        merge = self.request_json(
            "/api/merge/config",
            method="POST",
            files=[
                ("template", app.merge_service.DEFAULT_TEMPLATE),
                ("rules", app.merge_service.DEFAULT_RULES),
            ],
        )

        self.assertEqual(declaration["active"]["template"]["source"], "uploaded")
        self.assertEqual(merge["active"]["template"]["source"], "uploaded")
        self.assertTrue((app.TEMPLATE_DIR / "template.xlsx").is_file())
        self.assertTrue((app.RULES_DIR / "rules.xlsx").is_file())
        self.assertTrue((app.merge_service.TEMPLATE_DIR / "current.xlsx").is_file())
        self.assertTrue((app.merge_service.RULES_DIR / "current.xlsx").is_file())
        self.assertNotEqual(app.DATA_DIR, app.merge_service.DATA_DIR)

    def test_declaration_template_expands_beyond_eighteen_groups(self) -> None:
        lines = [
            {
                "itemNo": str(index),
                "hsCode": f"870899{index:04d}",
                "goodsName": f"测试商品 {index}",
                "quantity": index,
                "amount": index * 10,
                "currency": "USD",
                "brand": "测试品牌",
                "netWeight": index / 10,
            }
            for index in range(1, 26)
        ]
        preview = {
            "contractNo": "SP-25-GROUPS",
            "exportDateSerial": "",
            "consignee": "TEST CUSTOMER",
            "packageKind": "PALLET",
            "packageCount": 2,
            "grossWeight": 100,
            "netWeight": 90,
            "tradeTerm": "FOB",
            "originCountry": "中国",
            "destinationCountry": "意大利",
            "domesticSource": "苏州",
            "totals": {
                "quantity": sum(line["quantity"] for line in lines),
                "amount": sum(line["amount"] for line in lines),
                "currency": "USD",
            },
            "commodityLines": lines,
            "auditSamples": [],
        }
        output_path, _ = app.generate_workbook(preview)
        app.validate_xlsx_file(output_path, require_dimensions=True)
        workbook = load_workbook(output_path, data_only=False)
        worksheet = workbook["Sheet1"]
        self.assertEqual(worksheet["A18"].value, "1")
        self.assertEqual(worksheet["A90"].value, "25")
        self.assertEqual(worksheet["A93"].value, "Sub Total")
        workbook.close()

    @unittest.skipUnless(
        DECLARATION_INVOICE.is_file() and DECLARATION_PACKING.is_file(),
        "本地真实报关样例未包含在当前检出目录中",
    )
    def test_declaration_http_preview_generate_download_history_delete(self) -> None:
        preview = self.request_json(
            "/api/parse",
            method="POST",
            files=[
                ("documents", DECLARATION_INVOICE),
                ("documents", DECLARATION_PACKING),
            ],
        )
        self.assertEqual(preview["preview"]["contractNo"], "SP26000096")
        self.assertEqual(len(preview["preview"]["commodityLines"]), 15)

        generated = self.request_json(
            "/api/generate",
            method="POST",
            payload={
                "sessionId": preview["sessionId"],
                "historyId": preview["historyId"],
                "preview": preview["preview"],
            },
        )
        self.assert_xlsx(self.download(generated["downloadUrl"]))
        history = self.request_json("/api/history?limit=5")
        self.assertEqual(len(history["history"]), 1)
        self.assertTrue(history["history"][0]["outputName"])
        self.assert_xlsx(
            self.download(
                f"/api/history/{preview['historyId']}/download?kind=output"
            )
        )
        self.request_json(
            f"/api/history/{preview['historyId']}",
            method="DELETE",
        )
        self.assertEqual(self.request_json("/api/history?limit=5")["history"], [])

    def test_merge_http_preview_generate_download_history_delete(self) -> None:
        preview = self.request_json(
            "/api/merge/preview",
            method="POST",
            files=[("documents", self.merge_source)],
        )
        self.assertEqual(preview["preview"]["summary"]["sourceCount"], 1)
        self.assertEqual(preview["preview"]["summary"]["rowCount"], 1)

        generated = self.request_json(
            "/api/merge/generate",
            method="POST",
            payload={
                "sessionId": preview["sessionId"],
                "historyId": preview["historyId"],
                "confirmed": True,
            },
        )
        self.assert_xlsx(self.download(generated["downloadUrl"]))
        history = self.request_json("/api/merge/history?limit=5")
        self.assertEqual(len(history["history"]), 1)
        self.assertTrue(history["history"][0]["outputName"])
        self.assert_xlsx(
            self.download(
                f"/api/merge/history/{preview['historyId']}/download?kind=output"
            )
        )
        self.request_json(
            f"/api/merge/history/{preview['historyId']}",
            method="DELETE",
        )
        self.assertEqual(
            self.request_json("/api/merge/history?limit=5")["history"],
            [],
        )


if __name__ == "__main__":
    unittest.main()
